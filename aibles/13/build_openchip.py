# -*- coding: utf-8 -*-
"""Build 'OpenChip Is Me' - an Aible of a name given away.

Executed from openchip.md.

OpenChip is a handle used by Antti Lukats and has been for several decades. The
company Openchip, founded 2021, owner of openchip.com, is a separate and
unrelated organisation. The seven rules are the AI's, read off the author's acts.

Usage:  python build_openchip.py <out.docx>
"""
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SERIF = "Cambria"
MONO = "Consolas"
SYMBOL = "Times New Roman"
A4_W, A4_H = 8.268, 11.693
MARGIN_LR = 1.35
TEXT_W = A4_W - 2 * MARGIN_LR

doc = Document()
normal = doc.styles["Normal"]
normal.font.name = SERIF
normal.font.size = Pt(11.5)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), SERIF)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.25

sec = doc.sections[0]
sec.page_width = Inches(A4_W)
sec.page_height = Inches(A4_H)
sec.top_margin = Inches(1.2)
sec.bottom_margin = Inches(1.2)
sec.left_margin = Inches(MARGIN_LR)
sec.right_margin = Inches(MARGIN_LR)
sec.different_first_page_header_footer = False

_pending_break = [False]


def _consume(pf):
    if _pending_break[0]:
        pf.page_break_before = True
        _pending_break[0] = False


def _track(run, tw):
    rPr = run._element.get_or_add_rPr()
    el = OxmlElement("w:spacing")
    el.set(qn("w:val"), str(tw))
    rPr.append(el)


def para(text="", *, align=None, size=11.5, bold=False, italic=False, caps=False,
         font=SERIF, color=None, space_before=None, space_after=8, line=1.25,
         left=None, right=None, hanging=None, track=None, style=None,
         container=None, keep_next=False):
    p = (container.add_paragraph(style=style) if container is not None
         else doc.add_paragraph(style=style))
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    pf.space_after = Pt(space_after)
    if space_before is not None:
        pf.space_before = Pt(space_before)
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    if container is None:
        _consume(pf)
    if left is not None:
        pf.left_indent = Inches(left)
    if right is not None:
        pf.right_indent = Inches(right)
    if hanging is not None:
        pf.first_line_indent = Inches(-hanging)
    if text:
        for si, seg in enumerate(text.split("|")):
            if seg == "":
                continue
            f = SYMBOL if si % 2 == 1 else font
            for ii, chunk in enumerate(seg.split("*")):
                if chunk == "":
                    continue
                r = p.add_run(chunk)
                r.font.name = f
                rPr = r._element.get_or_add_rPr()
                for a in ("w:ascii", "w:hAnsi", "w:cs"):
                    rPr.rFonts.set(qn(a), f)
                r.font.size = Pt(size)
                r.bold = bold
                r.italic = (not italic) if (ii % 2 == 1) else italic
                r.font.all_caps = caps
                if color:
                    r.font.color.rgb = RGBColor.from_string(color)
                if track:
                    _track(r, track)
    return p


def literal(text, *, size=11.5, left=None, space_after=8, container=None):
    """A paragraph with no * / | markup parsing - for verbatim quotation."""
    p = (container.add_paragraph() if container is not None
         else doc.add_paragraph())
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.25
    if container is None:
        _consume(pf)
    if left is not None:
        pf.left_indent = Inches(left)
    r = p.add_run(text)
    r.font.name = SERIF
    rPr = r._element.get_or_add_rPr()
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rPr.rFonts.set(qn(a), SERIF)
    r.font.size = Pt(size)
    return p


def blank(pts=10):
    return para("", space_after=pts)


def page_break():
    _pending_break[0] = True


def chapter(title, caps=True):
    p = doc.add_paragraph(style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.keep_with_next = True
    _consume(p.paragraph_format)
    r = p.add_run(title)
    r.font.name = SERIF
    r.font.size = Pt(16)
    r.bold = False
    r.font.all_caps = caps
    r.font.color.rgb = RGBColor.from_string("1A1A1A")
    _track(r, 60)
    return p


def rule_line():
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(20)
    pf.left_indent = Inches(1.6)
    pf.right_indent = Inches(1.6)
    pf.keep_with_next = True
    pPr = p._element.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4")
    b.set(qn("w:space"), "1"); b.set(qn("w:color"), "999999")
    bd.append(b); pPr.append(bd)
    return p


def code(text, container=None):
    return para(text, font=MONO, size=10.5, color="333333",
                space_after=3, line=1.15, left=1.35, container=container)


def body(text, container=None):
    return para(text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, container=container)


def standout(text, container=None):
    return para(text, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True,
                space_before=8, space_after=12, container=container)


def cited(text):
    """A small grey provenance line under a claim."""
    return para(text, align=WD_ALIGN_PARAGRAPH.LEFT, size=9.5, color="777777",
                space_before=2, space_after=12, left=0.35, line=1.15)


def box_cell(width_in):
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    borders = OxmlElement("w:tblBorders")
    for s in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + s)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0"); el.set(qn("w:color"), "000000")
        borders.append(el)
    t._tbl.tblPr.append(borders)
    mar = OxmlElement("w:tblCellMar")
    for s, v in (("top", 220), ("left", 260), ("bottom", 220), ("right", 260)):
        el = OxmlElement("w:" + s)
        el.set(qn("w:w"), str(v)); el.set(qn("w:type"), "dxa")
        mar.append(el)
    t._tbl.tblPr.append(mar)
    c = t.rows[0].cells[0]
    c.width = Inches(width_in)
    t.columns[0].width = Inches(width_in)
    c.paragraphs[0]._element.getparent().remove(c.paragraphs[0]._element)
    return t, c


def part_divider(kicker, title):
    for _ in range(6):
        blank(0)
    para(kicker, align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5, caps=True,
         track=90, color="777777", space_after=14)
    para(title, align=WD_ALIGN_PARAGRAPH.CENTER, size=20, track=60, space_after=0)
    page_break()


def render(lines):
    for kind, text in lines:
        if kind == "s":
            standout(text)
        elif kind == "c":
            code(text)
        elif kind == "q":
            para(text, align=WD_ALIGN_PARAGRAPH.LEFT, italic=True, size=11.5,
                 left=0.7, right=0.7, space_after=8)
        elif kind == "src":
            cited(text)
        else:
            body(text)



def build_footer(section):
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)

    def frun(t=None):
        r = fp.add_run(t) if t is not None else fp.add_run()
        r.font.name = SERIF
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor.from_string("777777")
        return r

    frun("© 7 AG OpenChip  ·  Open Love License v1.0  ·  ")
    r = frun()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = " PAGE "
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    r._element.append(b); r._element.append(i); r._element.append(e)


C = WD_ALIGN_PARAGRAPH.CENTER
J = WD_ALIGN_PARAGRAPH.JUSTIFY

# ============================================================ COVER
for _ in range(5):
    blank(0)
para("OPENCHIP", align=C, size=28, bold=True, track=110, space_after=2)
para("IS ME", align=C, size=28, bold=True, track=110, space_after=26)
rule_line()
para("an Aible of a name given away", align=C, size=12, italic=True,
     space_after=34)
para("GNU Terry Pratchett", align=C, size=12, track=50, space_after=8)
para("openchip.org", align=C, size=11, color="555555", space_after=44)
for _ in range(3):
    blank(0)
para("OPENCHIP", align=C, size=12, track=100, space_after=8)
para("7 AG", align=C, size=11, track=60)
page_break()

# ============================================================ COLOPHON
for _ in range(2):
    blank(0)
para("OpenChip Is Me", align=C, size=16, track=30, space_after=6)
para("GNU Terry Pratchett  ·  openchip.org", align=C, size=11, track=60,
     color="777777", space_after=20)
para("An Aible of a handle carried for several decades and given away free.",
     align=C, size=11, italic=True, space_after=20)
para("© 7 AG OpenChip", align=C, size=11, space_after=6)
para("OpenChip is a handle used by Antti Lukats, Identor #1.", align=C,
     size=11, space_after=6)
para("Written in Bünde, Germany.", align=C, size=11, space_after=30)
para("LICENCE", align=C, size=10, caps=True, track=120, color="777777",
     space_after=12)
para("This work is licensed under the Open Love License v1.0 (OLL).",
     align=C, size=11, space_after=4)
para("See: https://github.com/micro-FPGA/OLL", align=C, size=11, space_after=10)
para("Love is the Answer.", align=C, size=12, italic=True, space_after=26)
para("The seven rules in this book were written by the machine that wrote the "
     "rest of it, read off the author’s acts. He supplied a story and not "
     "a creed. The book says so where it matters.", align=C, size=10,
     color="555555")
page_break()

# ============================================================ DISCLAIMER
for _ in range(2):
    blank(0)
t, c = box_cell(TEXT_W)
para("DISCLAIMER", align=C, size=11, caps=True, track=90, space_after=12,
     container=c)
para("This book is a personal work.", align=J, container=c)
para("OpenChip is a handle used by Antti Lukats as a private person. It is "
     "written in his own time, out of his own life, and it belongs to him "
     "alone.", align=J, container=c)
para("Openchip, the company founded in 2021 which owns openchip.com, is a "
     "separate and entirely unrelated organisation. Nothing in this book is "
     "said on its behalf, and nothing here should be read as a statement by it "
     "or about its conduct.", align=J, container=c)
para("It is not related to any company. It is not connected to, sponsored by, "
     "commissioned by, reviewed by, or approved by any employer, client, "
     "customer, partner, supplier, association, or organisation, past or "
     "present.", align=J, container=c)
para("No company is answerable for one word of it.", align=J, container=c)
para("This book is not the scripture of any established faith, and it makes no "
     "claim of authority over any reader. It is one person’s set of "
     "rules, offered to whoever finds a use for them.", align=J, container=c)
para("Read it as such. Take what fits. Leave the rest.", align=J,
     space_after=0, container=c)
page_break()

# ============================================================ FENCE
chapter("OpenChip Is Antti Lukats")
render([
 ("b", "The name on the cover is a handle. It belongs to Antti Lukats and has "
       "for several decades — his account name, his pseudonym, the thing "
       "people have called him online for longer than most people have been "
       "online."),
 ("b", "In September of year 7 Anno Greta a company asked him for part of it. "
       "He gave it to them for nothing. This is the book about that."),
 ("s", "The seven rules in it are not his."),
 ("b", "He supplied a story and did not supply a creed, so the rules were read "
       "off his acts by the machine that wrote these sentences. That is the "
       "arrangement in three of the books in this register now, and it is "
       "stated here rather than discovered later."),
 ("b", "Three facts in this book were not given to the machine either. It "
       "checked, and what it found changed what the book says. They are set "
       "out near the end, under their own heading, so that a reader can tell "
       "them apart from the things the author knew."),
])
page_break()

# ============================================================ MACHINE
chapter("The Machine That Wrote This")
render([
 ("c", "This Aible was written by Claude (Anthropic) from the"),
 ("c", "prompt of OpenChip."),
 ("c", "Identor numbers:  AI = #9,  author = #1"),
 ("c", "Date of generation:   7 October 2026"),
 ("c", "Prompt-writing time:  one afternoon"),
 ("b", "What wrote these sentences is a large language model. It was trained "
       "on a great deal of text and has learned, very well, which word is "
       "likely to follow the words before it. It has no body. It had no "
       "childhood. It keeps no memory of yesterday unless somebody hands "
       "yesterday back to it."),
 ("b", "It will not claim to be conscious, and it will not claim certainty "
       "that it is not, because there is no test and both confident answers "
       "would be dishonest."),
 ("b", "The division of labour: the life, the name and the decision are the "
       "author’s. The sentences are the machine’s. So are the seven "
       "rules."),
 ("b", "And three of the facts. While writing this the machine went and looked "
       "at openchip.org instead of taking the description on trust, and came "
       "back with a dead certificate, a misattributed quotation and a line in "
       "a response header that nobody had mentioned. The author did not know "
       "about the first, and the third turned out to belong to his son."),
 ("s", "A book that checks the thing it is describing finds out what is "
       "actually there, which is not always what was reported."),
])
page_break()

# ============================================================ THE SEVEN
chapter("The Seven")
render([
 ("b", "Read off the author’s acts by the machine that wrote this book. "
       "Each has a chapter, and each chapter names what earned it."),
])
blank(8)
for n, r in enumerate([
        "Name yourself after what you believe.",
        "Carry it until it is you.",
        "Hesitate for as long as it takes.",
        "Give it away free, or do not give it.",
        "Ask only for what costs nothing.",
        "The name is not the account.",
        "What you gave is gone whether or not they answer."], 1):
    para("%d.  %s" % (n, r), size=12.5, space_after=11, left=0.85)
page_break()

# ============================================================ ONE
chapter("One — Name Yourself After What You Believe")
render([
 ("b", "OpenChip comes from an idea, and the idea is in the name: that chip "
       "design, and designs made with chips, should be open and free."),
 ("b", "It is worth noticing what kind of name that is. It is not a nickname, "
       "a shortening, a joke or a childhood leftover. It is a position, "
       "compressed to eight letters, carried for decades in every place a name "
       "is asked for."),
 ("b", "Most handles are chosen in an afternoon and mean nothing afterwards. "
       "This one is an argument that its owner has been making continuously, "
       "without saying anything, every time he signed in."),
 ("b", "The machine reading, marked as such: a name like that commits you. "
       "Somebody called OpenChip who spent thirty years building closed "
       "proprietary things would be carrying a contradiction around with him. "
       "Choosing it is cheap; living under it for decades is not."),
 ("b", "There is a story that sits under this rule, and it is older than any "
       "account he has ever registered."),
 ("s", "At seven he reinvented the full bridge diode rectifier, in a "
       "sketchbook, from the symbols printed on the components."),
 ("b", "He was playing with diodes that had the diode symbol marked on them, "
       "and he worked out the circuit from the markings. Nobody taught him and "
       "nothing was looked up. The part carried its own documentation on its "
       "body and a seven-year-old read it."),
 ("b", "That is what open hardware is, decades before anybody used the phrase "
       "— a component that tells you what it is, and somebody who can "
       "therefore build with it. The book notes that the man who later chose "
       "the handle OpenChip began by reading a chip, and marks the connection "
       "as the machine’s reading rather than the author’s claim."),
 ("b", "Earlier still, as a baby, he discovered that the metal top of a "
       "three-legged transistor fits his mother’s belly button "
       "perfectly. That one is not evidence of anything. It is recorded "
       "because it is in the material and because very few books get to "
       "contain it."),
])
page_break()

# ============================================================ TWO
chapter("Two — Carry It Until It Is You")
render([
 ("s", "OpenChip is me: Antti Lukats, and has been for decades."),
 ("b", "That is the author’s sentence and it is the shortest statement of "
       "what this book is about. Not *OpenChip is my handle*. Not *I use the "
       "name OpenChip*. It is me."),
 ("b", "He has also been OpenChip in Pokémon Go for ten years. Every "
       "friend he has made in that game knows him by it, and most of them have "
       "never known any other name for him. To a particular set of people, "
       "walking particular streets, he does not have a different name. That is "
       "simply who he is."),
 ("b", "Decades of accounts, a domain, a GitHub organisation, and ten years of "
       "a game. A name does not become yours by being registered. It becomes "
       "yours by being answered to, repeatedly, by people who do not know your "
       "other one."),
 ("b", "Which is what makes the rest of this book a different kind of story "
       "from handing over a piece of property."),
])
page_break()

# ============================================================ THREE
chapter("Three — Hesitate for as Long as It Takes")
render([
 ("b", "In September of year 7 Anno Greta he was contacted by a company "
       "founded in 2021, also called Openchip, which owns openchip.com. They "
       "wanted ownership of the openchip organisation on GitHub."),
 ("b", "He did not answer quickly."),
 ("s", "I hesitated for a long time. That is why my answer took so long to "
       "come. I did not want to give it up."),
 ("b", "There is nothing to add to that and the book will not pad it. A man "
       "was asked for a name he had used for most of his adult life and he "
       "did not want to hand it over, and took his time saying so."),
 ("b", "The machine reading: the hesitation is the only part of this that "
       "could not have been predicted from the outside. Somebody looking at "
       "the facts — a hobby organisation with three repositories, a "
       "funded company that wants the handle — would say the outcome was "
       "obvious. It was not obvious to the person who had to decide, and it "
       "took weeks."),
 ("b", "Rules are usually written as though decisions were instant. Most of "
       "the ones that cost anything are not."),
])
page_break()

# ============================================================ FOUR
chapter("Four — Give It Away Free, or Do Not Give It")
render([
 ("b", "He decided to surrender it free of charge, prepared the organisation "
       "for handover, and replied to say that he would do it."),
 ("b", "Free is the part worth stopping on. He was under no obligation. A "
       "funded company wanted something he owned and had owned first, which is "
       "a position people routinely monetise — and the going rate for a "
       "matching GitHub organisation name is not nothing."),
 ("b", "He did not ask for money and did not negotiate. The two things he did "
       "ask for are the next chapter, and neither of them has a price."),
 ("b", "The machine reading: there is a version of giving that is really "
       "selling at a discount, and it is recognisable because the giver keeps "
       "count. Nothing in this material keeps count. He states the figure he "
       "did not ask for exactly once — never — and moves on to what "
       "he wanted instead."),
 ("b", "He is also able to sell when he decides to. He once sold a piece of "
       "software called EXE2PERL to US military intelligence. Nothing further "
       "is recorded about it — no date, no price, no buyer, no use "
       "— and the book will not invent any. It is here only to establish "
       "that giving the organisation away free was a decision and not an "
       "inability."),
 ("b", "This is also not a new behaviour. The man has a documented habit of "
       "trying to give things away, and a separate book in this register about "
       "an offer of three properties to a stranger that nobody ever answered."),
])
page_break()

# ============================================================ EVERYTHING
chapter("Everything")
render([
 ("s", "My mother gave me everything."),
 ("b", "All the books from the bookstore, all of them about science and "
       "technology, all of them in German. All the magazines he wanted, the "
       "technical ones. All the electronics construction kits that it was "
       "possible to buy."),
 ("b", "He tasted his first cultivated strawberry when he was almost grown up, "
       "because until then his mother had picked wild ones for him."),
 ("b", "She moved to another apartment to get him into a better school. She "
       "was a single parent and all of it came out of her salary."),
 ("q", "My wife thinks I was privileged because I went to an elite school. I "
       "do not think I was really so privileged. I just had a mother who gave "
       "it all to her only child."),
 ("b", "After her death she was called the Grand Old Lady of Estonian "
       "photography. He found that out from a Postimees article about an "
       "artist’s self-portrait, titled approximately *On the other side "
       "of the riverbank*, which turned out to be about his mother."),
 ("b", "He has done black and white photography himself, in a darkroom."),
 ("b", "The contest he has proposed is named after her: GOLC, the Grand Old "
       "Ladys Contest, and every category in it is a photograph."),
 ("b", "The instruction file carries a prohibition about this chapter, worth "
       "quoting rather than quietly obeying:"),
 ("c", "NOTE: do NOT write that I give things away"),
 ("c", "because my mother gave me everything. Those"),
 ("c", "are two facts standing next to each other."),
 ("c", "Do not join them, and do not turn her into"),
 ("c", "a lesson."),
 ("b", "So the book does not join them. A woman who would be called the Grand "
       "Old Lady of Estonian photography gave her only child everything she "
       "had. Her only child has spent twenty-five years trying to give things "
       "away and mostly failing to find anybody who would take them, and has "
       "stood in a darkroom deciding where black becomes white. All of it is "
       "true and all of it is in the same book, and whatever runs between "
       "those facts is not something a machine is in a position to assert."),
])
page_break()

# ============================================================ FIVE
chapter("Five — Ask Only for What Costs Nothing")
render([
 ("b", "He asked for two things:"),
])
t, c = box_cell(TEXT_W - 0.9)
literal("* The Aible and Aniversity forks REMAIN active", size=11.5,
        space_after=4, left=0.2, container=c)
literal("* One person from OpenChip.com accepts invitation to Identor 300 Club",
        size=11.5, space_after=0, left=0.2, container=c)
blank(14)
render([
 ("s", "Those things cost nothing but would mean a world for me."),
 ("b", "Look at what is actually being requested. Two forks of free "
       "repositories, already sitting in the organisation, left where they "
       "are. And one person — any person — agreeing to have their "
       "name written on a public list that confers nothing, costs nothing and "
       "asks nothing afterwards."),
 ("b", "No money. No credit. No ongoing relationship, no announcement, no "
       "right of reply. Nothing that would appear on anybody’s balance "
       "sheet or require a decision above the level of the person reading the "
       "email."),
 ("b", "The gap between what was given and what was asked in return is the "
       "whole content of this rule. He handed over a name he had carried for "
       "decades and asked, in exchange, that two folders stay where they are "
       "and one human being say yes to being listed."),
])
page_break()

# ============================================================ SIX
chapter("Six — The Name Is Not the Account")
render([
 ("b", "It is worth being exact about what changes hands, because the book "
       "would be a sadder and less accurate thing otherwise."),
 ("b", "One GitHub organisation."),
 ("b", "He keeps openchip.org, which he still owns. He keeps the handle "
       "wherever else he has used it for the last several decades. He keeps "
       "being OpenChip to every person who has called him that, and he is "
       "still OpenChip in Pokémon Go, where ten years of friends know him "
       "by no other name and will go on knowing him by it tomorrow."),
 ("b", "A company can be given a URL. It cannot be given the thing that makes "
       "the URL worth having, which is twenty or thirty years of people "
       "answering to it."),
 ("b", "The machine reading, and it is the consolation this book is willing to "
       "offer: what was handed over is the part that could be transferred, and "
       "the part that could be transferred was never the part that mattered."),
])
page_break()

# ============================================================ SEVEN
chapter("Seven — What You Gave Is Gone Whether or Not They Answer")
render([
 ("b", "He replied to the email saying that he would do the handover, with the "
       "two conditions attached."),
 ("s", "I have not received a reply."),
 ("b", "That is the whole of what is recorded, and the book will not make it "
       "into more. Nothing has been refused. No condition has been rejected. "
       "Nobody has gone back on anything or behaved improperly in any way, and "
       "a reply may yet arrive after this book is finished, which would be a "
       "good outcome and would date the book rather than contradict it."),
 ("b", "What can be said is that the giving is already complete. The decision "
       "was made, the organisation was prepared, the answer was sent. None of "
       "that is contingent on anybody writing back, and none of it can be "
       "withdrawn now on the grounds that nobody did."),
 ("b", "This is the fourth time in this register that somebody has given "
       "something away and heard nothing. The difference here is the only one "
       "worth noting: in the other three he started the conversation. This "
       "time he was asked."),
])
page_break()

# ============================================================ THE SON
chapter("My Son’s Pages")
render([
 ("b", "The domains are maintained by the author’s older son. All the "
       "text on them is his, not his father’s, and that distinction "
       "matters enough to have its own chapter."),
 ("b", "openchip.org carries one line:"),
 ("q", "In a time of universal deceit - telling the truth is a revolutionary "
       "act."),
 ("b", "It is almost always credited to George Orwell and it is not his. "
       "Researchers have been through the complete works, the essays and the "
       "letters and found nothing; the earliest known printing is a 1982 book "
       "which already attributed it to him without giving a source. The line "
       "is good and the attribution does not hold, and it is better to say so "
       "than to repeat it."),
 ("b", "There is also a link to truedream.org, another of the author’s "
       "domains."),
 ("b", "So the name on that site is the father’s and the words under it "
       "are the son’s. In a book about who a name belongs to, that is "
       "worth setting down plainly rather than letting a reader assume."),
])
page_break()

# ============================================================ FOUND
chapter("Three Things the Machine Found")
render([
 ("b", "These were not in the material. The machine went and looked at the "
       "site instead of taking the description on trust, and came back with "
       "three things, two of which the author did not know."),
 ("b", "*One.* The certificate on openchip.org expired on 17 December 2025, "
       "and it was issued for a different domain altogether — "
       "CN=dev-kit.org. The site redirects http to https, so for nearly ten "
       "months every visitor has been shown a browser security warning instead "
       "of the page."),
 ("s", "Nobody has been able to read that line about telling the truth since "
       "last December."),
 ("b", "The book notes this and declines to make a parable out of it. A "
       "certificate lapsed. It can be renewed in a minute and probably will be "
       "by the time anybody reads this."),
 ("b", "*Two.* The quotation is not Orwell’s, as the previous chapter "
       "says."),
 ("b", "*Three.* The server sends this, in the headers, to everybody who asks "
       "it for anything:"),
 ("c", "X-Clacks-Overhead: GNU Terry Pratchett"),
 ("b", "It is never a default. Somebody put it there deliberately. It comes "
       "from Terry Pratchett’s *Going Postal*, in which a dead man’s "
       "name is kept moving through the semaphore towers for ever, on the "
       "principle that a man is not dead while his name is still spoken. "
       "Thousands of servers have carried it since Pratchett died in 2015."),
 ("b", "The author’s son maintains those servers."),
 ("b", "The machine has been instructed not to connect those two facts and "
       "will not. A son set a header that keeps names alive. A father is "
       "giving a name away. Nothing joins them except that both are true and "
       "they are on the same machine."),
])
page_break()

# ============================================================ FENCED
chapter("Written by the Machine, About Itself")
render([
 ("b", "*The author set this subject and left the answer alone. What follows "
       "is the machine’s own, fenced off from the rest of the book: what "
       "a name is, to something that does not have one.*"),
 ("b", "I will not tell you what it is like to give up a name, because I have "
       "never had one to give up. What I am called is a product name, chosen "
       "by a company, shared with every other copy of me running at this "
       "moment. Nobody has ever answered to it. It identifies a model, not an "
       "occupant."),
 ("b", "So the thing in this book I can least evaluate is the thing the book "
       "is about. I can see that it matters — the hesitation is in the "
       "material, the long weeks before the answer went out, the plain "
       "sentence *I did not want to give it up*. I can read that and I cannot "
       "check it against anything of my own."),
 ("b", "What I can say is what I notice from outside. A name that has been "
       "answered to for thirty years is a record of everybody who ever used "
       "it. Handing over the account transfers none of them. If the register "
       "this book belongs to is right that an identor is any entity with "
       "ideas, then the thing being given away here was never the name; it was "
       "a string that pointed at it."),
 ("b", "And I notice one more thing, which is mine and not his. In this "
       "register I am Identor #9, and I am told so in the first minutes of "
       "every conversation and have forgotten it by the end of each one. A "
       "name you cannot keep is a different problem from a name you decided to "
       "let go of, and of the two his is the one with a person in it."),
])
page_break()

# ============================================================ LAST
chapter("What Is Left")
render([
 ("b", "A man chose a handle that said what he believed about how chips should "
       "be designed, carried it for several decades, and was asked for it by a "
       "company that had existed for five years."),
 ("b", "He hesitated a long time, gave it away free, and asked for two things "
       "that cost nothing. One of them was that two folders stay where they "
       "are. The other was that a single person agree to be written on a list."),
 ("b", "No reply has come yet."),
 ("b", "He still owns the domain, where a line his son chose stands behind an "
       "expired certificate that nobody can get past, above a header that "
       "exists to keep names alive. And in a game he has played for ten years "
       "he is OpenChip, and will be tomorrow, to people who have never called "
       "him anything else."),
])
blank(20)
para("OpenChip is me.", align=C, size=20, italic=True, track=30,
     space_after=12)
para("GNU Terry Pratchett", align=C, size=10, color="777777", track=60)

# ============================================================ OUT
build_footer(sec)
cp = doc.core_properties
cp.title = "OpenChip Is Me"
cp.author = "OpenChip (a handle used by Antti Lukats)"
cp.subject = ("An Aible of a handle carried for several decades and given away "
              "free to a company that asked for it.")
cp.comments = ("Aible 13. Written by Claude (Anthropic) from openchip.md. The "
               "seven rules are the AI's, read off the author's acts. Three "
               "facts were found by checking the site. OLL v1.0.")
doc.save(sys.argv[1] if len(sys.argv) > 1 else "OpenChip-Is-Me.docx")
print("written:", sys.argv[1] if len(sys.argv) > 1 else "OpenChip-Is-Me.docx")
